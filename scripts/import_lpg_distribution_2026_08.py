"""Impor laporan penyaluran LPG Agustus 2026 dan perkaya master pangkalan."""
from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "Dokumen Pangkalan" / "Penyaluran"
MASTER_PATH = ROOT / "assets" / "data" / "lpg-pangkalan.json"
AGENTS_PATH = ROOT / "assets" / "data" / "lpg-agents.json"
OUTPUT = ROOT / "assets" / "data" / "lpg-distribution-2026-08.json"
POLICY_OUTPUT = ROOT / "assets" / "data" / "lpg-policy-documents.json"

AGENTS = {
    "PT. KAKA MIGAS UTAMA": {"id": "AG-008", "soldTo": "1015166"},
    "PT WAHYU DWI KENCANA MANDIRI": {"id": "AG-005", "soldTo": "1070992"},
    "PT. HAMISA SUKRAH MULYA": {"id": "AG-002", "soldTo": "818234"},
}

REPORT_NAME_OVERRIDES = {
    # Koreksi pembacaan baris terbungkus pada layout PDF dan perubahan nama
    # yang dapat dipastikan melalui ID registrasi laporan Pertamina.
    "791254818234066": "ANITA",
    "791252818234045": "ANUGRAH PRATAMA PUTRA",
    "791271818234002": "JEMI LABIRO HAMZAH",
    "791256818234064": "NURASIA",
    "7912540818234001": "NURASIA",
    "791254818234072": "RUSNI SORONG",
    "7912541070992040": "HJ NADIRA 77",
    "7912121070992050": "PANGKALAN ADY JAYA MANDIRI",
    "7912131070992017": "RIECHA",
    "7912151070992034": "RIZKY TABUNG GAS",
}

MASTER_TARGET_OVERRIDES = {
    "791271818234002": "PG-000310",  # JEMI LABIRO -> JEMI LABIRO HAMZAH
    "791256818234064": "PG-000558",  # NURASIA, Tiroang
    "7912540818234001": "PG-000197", # NURASIA, Lembang
    "791254818234072": "PG-000181",  # RUSNI -> RUSNI SORONG
}


def norm(value):
    text = unicodedata.normalize("NFKD", str(value or "").upper())
    return re.sub(r"[^A-Z0-9]+", "", text)


def record(agent, registration_id, name, allocation, daily, normal, facultative, grand, source, **extra):
    assert len(daily) == 31
    assert sum(daily) == grand, (name, sum(daily), grand)
    assert normal + facultative == grand, (name, normal, facultative, grand)
    return {
        "period": "2026-08", "unit": "tabung", "agentId": AGENTS[agent]["id"],
        "agentName": agent, "soldTo": AGENTS[agent]["soldTo"],
        "registrationId": str(registration_id), "pangkalanName": str(name).strip(),
        "allocation": int(allocation),
        "dailyDistribution": {f"2026-08-{day:02d}": int(value) for day, value in enumerate(daily, 1)},
        "totalNormal": int(normal), "totalFacultative": int(facultative),
        "grandTotal": int(grand), "realizationPct": round(grand / allocation * 100, 2) if allocation else None,
        "sourceDocument": source, "sourceStatus": "AGENT_REPORTED_PERTAMINA_SYSTEM",
        **extra,
    }


def parse_excel():
    path = DOCS / "Penyaluran Agustus 2026.xlsx"
    ws = load_workbook(path, data_only=True).active
    rows = []
    for values in ws.iter_rows(min_row=2, max_row=47, values_only=True):
        rows.append(record(
            "PT. KAKA MIGAS UTAMA", values[0], values[1], values[7], list(values[8:39]),
            values[39], values[40], values[41], path.name,
            province=values[2], regency=values[3], district=values[4], village=values[5], address=values[6],
        ))
    for row in rows:
        row["pangkalanName"] = REPORT_NAME_OVERRIDES.get(row["registrationId"], row["pangkalanName"])
    return rows


def parse_pdf(path, agent):
    rows = []
    ignored_continuations = {"ID REGISTRASI", "NAMA", "PANGKALAN", "ALOKASI", "TOTAL", "NORMAL", "FAKULTATIF", "GRAND"}
    for page in PdfReader(str(path)).pages:
        lines = (page.extract_text(extraction_mode="layout") or "").splitlines()
        page_rows = []
        for line in lines:
            match = re.match(r"^\s*(791\d{12,13})\s+(.+)$", line)
            if match:
                registration_id, remainder = match.groups()
                tokens = list(re.finditer(r"(?<![\w.])\d{1,3}(?:,\d{3})*(?![\w.])", remainder))
                if len(tokens) < 35:
                    continue
                selected = tokens[-35:]
                numbers = [int(token.group().replace(",", "")) for token in selected]
                name = remainder[:selected[0].start()].strip()
                name = REPORT_NAME_OVERRIDES.get(registration_id, name)
                row = record(agent, registration_id, name, numbers[0], numbers[1:32], numbers[32], numbers[33], numbers[34], path.name)
                rows.append(row); page_rows.append(row)
                continue
            stripped = line.strip()
            if not page_rows or not stripped or stripped.upper() in ignored_continuations:
                continue
            # Nama panjang dicetak sebagai baris lanjutan tepat setelah baris data.
            if (re.fullmatch(r"[A-Za-zÀ-ÿ.,' ]{2,40}", stripped)
                    and not any(word in stripped.upper() for word in ("PAGE ", "PERTAMINA", "LAPORAN", "NAMA AGEN", "ALAMAT", "EMAIL", "WILAYAH", "MENGETAHUI", "PENERIMA", "JABATAN"))):
                current = page_rows[-1]
                if len(current["pangkalanName"].split()) <= 3:
                    current["pangkalanName"] = f'{current["pangkalanName"]} {stripped}'.strip()
    for row in rows:
        row["pangkalanName"] = REPORT_NAME_OVERRIDES.get(row["registrationId"], row["pangkalanName"])
    return rows


def attach_master(records, master):
    by_registration = {str(x.get("registrationId")): x for x in master if x.get("registrationId")}
    by_agent_name = defaultdict(list)
    by_id = {x["id"]: x for x in master}
    for item in master:
        by_agent_name[(item.get("agentId"), norm(item.get("name")))].append(item)
    for row in records:
        target = by_registration.get(row["registrationId"]) or by_id.get(MASTER_TARGET_OVERRIDES.get(row["registrationId"]))
        if not target:
            choices = by_agent_name[(row["agentId"], norm(row["pangkalanName"]))]
            target = choices[0] if len(choices) == 1 else None
        row["pangkalanId"] = target.get("id") if target else None
        row["masterMatchStatus"] = "MATCHED_REGISTRATION_ID" if target and target.get("registrationId") == row["registrationId"] else "MATCHED_UNIQUE_NAME" if target else "UNMATCHED"
        if target:
            target["registrationId"] = row["registrationId"]
            target["monthlyAllocation"] = row["allocation"]
            target["allocationUnit"] = "tabung"
            target["latestDistributionPeriod"] = "2026-08"
            target["latestDistributionNormal"] = row["totalNormal"]
            target["latestDistributionFacultative"] = row["totalFacultative"]
            target["latestDistributionTotal"] = row["grandTotal"]
            target["latestDistributionUnit"] = "tabung"
            target["latestDistributionSource"] = row["sourceDocument"]
    return by_registration, by_agent_name


def main():
    master = json.loads(MASTER_PATH.read_text(encoding="utf-8"))
    agents = json.loads(AGENTS_PATH.read_text(encoding="utf-8"))
    # Pulihkan kode agen yang salah pada entri baru hasil impor dokumen sebelumnya.
    for item in master:
        expected = AGENTS.get(item.get("agentName"))
        if expected:
            item["agentId"] = expected["id"]

    records = parse_excel()
    records += parse_pdf(DOCS / "Penyaluran Harian KABUPATEN PINRANG 2026-08 (1).pdf", "PT WAHYU DWI KENCANA MANDIRI")
    records += parse_pdf(DOCS / "Penyaluran Harian KABUPATEN PINRANG 2026-08 (3).pdf", "PT. HAMISA SUKRAH MULYA")
    assert len(records) == 159, len(records)
    assert len({(x["agentId"], x["registrationId"]) for x in records}) == 159

    attach_master(records, master)
    # Pembaruan identitas yang ditopang ID registrasi laporan Pertamina.
    rename_by_registration = {"791271818234002": "JEMI LABIRO HAMZAH", "791254818234072": "RUSNI SORONG"}
    for item in master:
        if item.get("registrationId") in rename_by_registration:
            item.setdefault("previousDistributionDocumentName", item["name"])
            item["name"] = rename_by_registration[item["registrationId"]]
            item["normalizedName"] = norm(item["name"])
        if item.get("registrationId") == "791253818234058":
            item["identityReviewRequired"] = True
            item["identityReviewReason"] = "Nama '74.91270' menyerupai kode SPBU dan harus dikonfirmasi kepada agen."

    # HJ.MINA adalah satu-satunya registrasi laporan Hamisa yang belum ada di master.
    hj_mina = next(x for x in records if x["registrationId"] == "791215818234028")
    if not any(x.get("registrationId") == hj_mina["registrationId"] for x in master):
        next_id = max(int(x["id"].split("-")[1]) for x in master) + 1
        new_item = {
            "id": f"PG-{next_id:06d}", "agentId": "AG-002", "agentName": "PT. HAMISA SUKRAH MULYA",
            "registrationId": hj_mina["registrationId"], "name": "HJ.MINA", "normalizedName": "HJMINA",
            "kecamatan": None, "kecamatanSource": None, "desaKelurahan": None, "desaKelurahanSource": None,
            "address": None, "latitude": None, "longitude": None, "googleMapsUrl": None,
            "pointType": "UNMAPPED", "gpsVerified": False, "verificationStatus": "identity-review",
            "status": "candidate", "isDeleted": False, "sourceType": "PERTAMINA_DISTRIBUTION_REPORT",
            "sourceDate": "2026-09-04", "sourceUrl": f'Dokumen Pangkalan/Penyaluran/{hj_mina["sourceDocument"]}',
            "sourceLine": None, "monthlyAllocation": hj_mina["allocation"], "allocationUnit": "tabung",
            "latestDistributionPeriod": "2026-08", "latestDistributionNormal": hj_mina["totalNormal"],
            "latestDistributionFacultative": hj_mina["totalFacultative"], "latestDistributionTotal": hj_mina["grandTotal"],
            "latestDistributionUnit": "tabung", "latestDistributionSource": hj_mina["sourceDocument"],
            "identityReviewRequired": True,
            "identityReviewReason": "Registrasi ditemukan pada laporan Pertamina tetapi identitas wilayah/alamat belum tersedia.",
            "dataVersion": "2026-09-30-distribution-import",
        }
        master.append(new_item)
        hj_mina["pangkalanId"] = new_item["id"]
        hj_mina["masterMatchStatus"] = "ADDED_CANDIDATE_FROM_DISTRIBUTION_REPORT"

    summaries = []
    for agent_name, meta in AGENTS.items():
        subset = [x for x in records if x["agentName"] == agent_name]
        summaries.append({
            "agentId": meta["id"], "agentName": agent_name, "soldTo": meta["soldTo"], "pangkalanCount": len(subset),
            "allocation": sum(x["allocation"] for x in subset), "totalNormal": sum(x["totalNormal"] for x in subset),
            "totalFacultative": sum(x["totalFacultative"] for x in subset), "grandTotal": sum(x["grandTotal"] for x in subset),
            "unit": "tabung", "zeroDistributionCount": sum(x["grandTotal"] == 0 for x in subset),
        })
    payload = {
        "schemaVersion": 1, "period": "2026-08", "unit": "tabung", "status": "agent-reported",
        "generatedAt": "2026-09-30T00:00:00+08:00", "recordCount": len(records),
        "sourceNote": "Angka ditetapkan sebagai jumlah tabung sesuai konfirmasi pengguna.",
        "summaries": summaries, "records": records,
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    MASTER_PATH.write_text(json.dumps(master, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    agent_contacts = {
        "AG-005": {"soldTo": "1070992", "distributionReportEmail": "wahyudwikencana.mandiri@gmail.com"},
        "AG-002": {"soldTo": "818234", "distributionReportEmail": "pthamisasukrahmulya@yahoo.co.id"},
        "AG-008": {"soldTo": "1015166", "distributionReportEmail": None},
    }
    counts = Counter(x["agentId"] for x in master if not x.get("isDeleted"))
    for agent in agents:
        if agent["id"] not in agent_contacts:
            continue
        agent.update(agent_contacts[agent["id"]])
        agent["linkedPangkalanCount"] = counts[agent["id"]]
        agent["latestDistributionPeriod"] = "2026-08"
        agent["distributionDataUnit"] = "tabung"
        agent["distributionDataSource"] = next(x["sourceDocument"] for x in records if x["agentId"] == agent["id"])
        agent["relationshipStatus"] = "Terhubung ke master pangkalan dan laporan penyaluran Agustus 2026"
    AGENTS_PATH.write_text(json.dumps(agents, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    policies = {
        "schemaVersion": 1,
        "documents": [
            {
                "id": "LPG-CIRCULAR-2026-08-14-DRAFT", "type": "CIRCULAR_DRAFT", "status": "DRAFT_UNNUMBERED",
                "title": "Penggunaan Gas LPG Subsidi (3 Kg) Tepat Sasaran", "documentDate": "2026-08-14",
                "number": None, "unit": "tabung", "publicationAllowed": False,
                "sourceDocument": "SURAT EDARAN BUPATI PINRANG ( LPG 3 KG ) (1).pdf",
                "reviewNotes": ["Nomor surat kosong", "Metadata PDF tidak sesuai isi", "Dasar hukum dan redaksi perlu verifikasi Bagian Hukum"],
            },
            {
                "id": "LPG-QUOTA-PROPOSAL-2026-09-03", "type": "EMERGENCY_QUOTA_PROPOSAL", "status": "PROPOSED_NOT_APPROVED",
                "title": "Permohonan Alokasi Tambahan Fakultatif Darurat LPG 3 Kg Bersubsidi", "documentDate": "2026-09-03",
                "number": None, "requestedQuantity": 5040, "unit": "tabung", "requestedLoads": 9,
                "quantityPerLoad": 560, "districtCount": 12, "statedHetRupiah": 18500,
                "publicationAllowed": False, "sourceDocument": "Usulan Penambahan Kuota LPG  3 KG BERSUBSIDI.pdf",
                "reviewNotes": ["Nomor surat kosong", "Lampiran peta klaster tidak terdapat dalam PDF", "Belum ada bukti persetujuan atau realisasi Pertamina"],
            },
        ],
    }
    POLICY_OUTPUT.write_text(json.dumps(policies, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"records": len(records), "master": len(master), "summaries": summaries, "matchStatus": Counter(x["masterMatchStatus"] for x in records)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
