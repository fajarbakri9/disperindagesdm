"""Update master pangkalan LPG dari dokumen agen di folder Dokumen Pangkalan.

Dokumen diperlakukan sebagai sumber administratif terbaru, bukan bukti verifikasi
GPS lapangan. Script membuat laporan audit JSON/CSV dan mempertahankan ID lama
untuk baris yang dapat dipadankan secara meyakinkan.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import unicodedata
from copy import deepcopy
from difflib import SequenceMatcher
from pathlib import Path
from urllib.parse import quote_plus

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "assets" / "data" / "lpg-pangkalan.json"
DOCS = ROOT / "Dokumen Pangkalan"
REPORT_JSON = DOCS / "HASIL_ANALISIS_UPDATE_PANGKALAN.json"
REPORT_CSV = DOCS / "HASIL_ANALISIS_UPDATE_PANGKALAN.csv"

AGENTS = {
    "nur": ("AG-004", "PT. NURCAHAYA ENERGI ABADI"),
    "wahyu": ("AG-005", "PT WAHYU DWI KENCANA MANDIRI"),
    "kmu": ("AG-008", "PT. KAKA MIGAS UTAMA"),
}

DISTRICTS = [
    "WATANG SAWITO", "MATIRRO SOMPE", "MATTIRO BULU", "BATU LAPPA",
    "PATAMPANUA", "PALETEANG", "DUAMPANUA", "LANRISANG", "LANSIRANG",
    "TIROANG", "LEMBANG", "SUPPA", "CEMPA",
]
VILLAGE_PREFIXES = ("DESA ", "KELURAHAN ")


def norm(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").upper())
    return re.sub(r"[^A-Z0-9]+", "", text)


def title_district(value: str) -> str:
    fixes = {"BATU LAPPA": "Batulappa", "LANSIRANG": "Lanrisang"}
    return fixes.get(value.upper(), value.title())


def clean_village(value: str) -> tuple[str, str]:
    source = value.strip().upper()
    bare = source
    for prefix in VILLAGE_PREFIXES:
        if bare.startswith(prefix):
            bare = bare[len(prefix):]
            break
    return bare.title(), source


def pdf_lines(path: Path) -> list[str]:
    text = "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
    return [re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip()]


def split_district(rest: str) -> tuple[str, str]:
    for district in sorted(DISTRICTS, key=len, reverse=True):
        if rest.startswith(district + " "):
            return district, rest[len(district):].strip()
    raise ValueError(f"Kecamatan tidak dikenali: {rest}")


def parse_nur(path: Path) -> list[dict]:
    rows = []
    marker = "PT. NURCAHAYA ENERGI ABADI PANGKALAN LPG PSO "
    rx_end = re.compile(r"\s+(-?\d{6,7})\s+(1\d{8})$")
    for line in pdf_lines(path):
        m_no = re.match(r"^(\d+)\s+", line)
        if not m_no or marker not in line:
            continue
        body = line.split(marker, 1)[1]
        reg, body = body.split(" ", 1)
        name, body = body.split(" SULAWESI SELATAN KABUPATEN PINRANG ", 1)
        district, body = split_district(body)
        end = rx_end.search(body)
        if not end:
            raise ValueError(f"Koordinat Nurcahaya tidak terbaca: {line}")
        lat, lon = int(end.group(1)) / 1_000_000, int(end.group(2)) / 1_000_000
        middle = body[:end.start()].strip()
        # Batas kelurahan ditentukan oleh RW/kode pos; untuk baris tanpa RW,
        # kode pos tetap menjadi pemisah yang stabil.
        z = re.search(r"\s+(?:RW-\d+\s+)?(91\d{3})\s+", middle)
        if not z:
            raise ValueError(f"Kode pos Nurcahaya tidak terbaca: {line}")
        village, address = middle[:z.start()].strip(), middle[z.end():].strip()
        rows.append(make_row("nur", reg, name, district, village, address, lat, lon, path.name, int(m_no.group(1))))
    return rows


def parse_wahyu(path: Path) -> list[dict]:
    rows = []
    rx = re.compile(r"^(791\d+)\s+(.+?)\s+KABUPATEN PINRANG\s+(.+?)\s+(-?\d+\.\d+)\s+(1\d+\.\d+)$")
    village_names = set()
    # Master administratif membantu mencari batas kecamatan/kelurahan yang
    # tidak memiliki label DESA/KELURAHAN di PDF.
    admin = json.loads((ROOT / "assets/data/pinrang-administrative-master.json").read_text(encoding="utf-8"))
    def walk(x):
        if isinstance(x, dict):
            for k, v in x.items():
                if str(k).lower() in {"desa", "kelurahan", "desakelurahan", "name", "nama"} and isinstance(v, str):
                    village_names.add(v.upper())
                walk(v)
        elif isinstance(x, list):
            for v in x: walk(v)
    walk(admin)
    # Nama yang muncul di dokumen tetapi dapat berbeda dari master administratif.
    village_names.update({"SAWITO","MACCORAWALIE","JAYA","BENTENG SAWITO","MAMMINASAE","PENRANG","SIPARAPPE","SIPATOKKONG","BENTENGNGE","TATAE","MALIMPUNG","MATTIROADE","SIWOLONG POLONG","CEMPA","MATUNRU TUNRUE","ALITTA","MANARANG","TELLUMPANUA","MACINNAE","SIKKUALE","PEKKABATA","TADDOKONG","LALENG BATA","MATTONGANG TONGANG","PAKKE","PAKKIE","MARAWI","PATOBONG"})
    for line in pdf_lines(path):
        m = rx.match(line)
        if not m:
            continue
        reg, name, body, lat, lon = m.groups()
        district, body = split_district(body)
        village = next((v for v in sorted(village_names, key=len, reverse=True) if body.startswith(v + " ")), None)
        if not village:
            raise ValueError(f"Kelurahan Wahyu tidak terbaca: {line}")
        address = body[len(village):].strip()
        rows.append(make_row("wahyu", reg, name, district, village, address, float(lat), float(lon), path.name, len(rows)+1))
    return rows


def parse_kmu(path: Path) -> list[dict]:
    ws = load_workbook(path, data_only=True).active
    rows = []
    for n, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 1):
        _, reg, name, district, village, address, lat, lon = values
        if not reg:
            continue
        rows.append(make_row("kmu", str(reg), str(name), str(district), str(village), str(address), float(lat), float(lon), path.name, n))
    return rows


def make_row(agent_key, reg, name, district, village, address, lat, lon, source, source_row):
    agent_id, agent_name = AGENTS[agent_key]
    village_title, village_source = clean_village(village)
    return {
        "agentKey": agent_key, "agentId": agent_id, "agentName": agent_name,
        "registrationId": str(reg), "name": name.strip(), "normalizedName": norm(name),
        "kecamatan": title_district(district), "kecamatanSource": district.upper(),
        "desaKelurahan": village_title, "desaKelurahanSource": village_source,
        "address": address.strip(), "latitude": lat, "longitude": lon,
        "documentSource": source, "documentRow": source_row,
    }


def similarity(old: dict, new: dict) -> float:
    name = SequenceMatcher(None, norm(old.get("name")), norm(new["name"])).ratio()
    village = SequenceMatcher(None, norm(old.get("desaKelurahanSource")), norm(new["desaKelurahanSource"])).ratio()
    district = 1.0 if norm(old.get("kecamatanSource")) == norm(new["kecamatanSource"]) else 0.0
    address = SequenceMatcher(None, norm(old.get("address")), norm(new["address"])).ratio()
    return 0.68 * name + 0.14 * village + 0.10 * district + 0.08 * address


def match_rows(master: list[dict], incoming: list[dict]):
    used = set()
    matches, additions = [], []
    for new in incoming:
        candidates = [(similarity(old, new), i, old) for i, old in enumerate(master)
                      if i not in used and old.get("agentName") == new["agentName"]]
        candidates.sort(reverse=True, key=lambda x: x[0])
        score, i, old = candidates[0] if candidates else (0, -1, None)
        old_name = norm(old.get("name")) if old is not None else ""
        exact_name = old is not None and old_name == new["normalizedName"]
        # Dokumen baru kadang menambah/menghapus penanda lokasi atau nomor pada
        # nama (contoh: ISMAIL -> ISMAIL TEPPO 4017, ALIM MADALLO -> ALIM).
        # Terima relasi substring hanya bila kecamatannya juga sama.
        renamed_same_district = bool(old is not None and min(len(old_name), len(new["normalizedName"])) >= 4
                                     and (old_name in new["normalizedName"] or new["normalizedName"] in old_name)
                                     and norm(old.get("kecamatanSource")) == norm(new["kecamatanSource"]))
        if old is not None and (score >= 0.72 or exact_name or renamed_same_district):
            used.add(i)
            matches.append((i, old, new, score))
        else:
            additions.append(new)
    return matches, additions


def update_record(old: dict, new: dict) -> dict:
    out = deepcopy(old)
    old_snapshot = {k: old.get(k) for k in ("name", "kecamatan", "desaKelurahan", "address", "latitude", "longitude")}
    for key in ("agentId", "agentName", "name", "normalizedName", "kecamatan", "kecamatanSource", "desaKelurahan", "desaKelurahanSource", "address", "latitude", "longitude"):
        out[key] = new[key]
    out.update({
        "registrationId": new["registrationId"],
        "googleMapsUrl": f"https://www.google.com/maps?q={new['latitude']:.7f},{new['longitude']:.7f}",
        "googleMapsSearchUrl": "https://www.google.com/maps/search/?api=1&query=" + quote_plus(f"{new['latitude']},{new['longitude']}"),
        "plusCode": None, "listingName": None,
        "locationPrecisionLevel": "B", "locationMatchStatus": "DOCUMENT_COORDINATE",
        "locationMethod": "AGENT_DOCUMENT_COORDINATE", "locationConfidence": "HIGH",
        "pointType": "DOCUMENT_COORDINATE", "gpsVerified": False,
        "locationVerification": "Koordinat tercantum dalam dokumen agen; belum dinyatakan sebagai hasil verifikasi GPS lapangan oleh Disperindag ESDM.",
        "verificationStatus": "document-sourced", "geolocationSource": new["documentSource"],
        "coordinateCorrectionNote": "Menggantikan titik administratif/koordinat lama berdasarkan dokumen agen terbaru.",
        "dataVersion": "2026-09-30-document-update",
        "documentSource": new["documentSource"], "documentRow": new["documentRow"],
    })
    # Simpan baseline pertama; eksekusi ulang tidak boleh menghapus jejak nilai
    # yang benar-benar mendahului pembaruan dokumen.
    out.setdefault("previousDocumentUpdateValue", old_snapshot)
    out["sourceOriginal"] = {**(old.get("sourceOriginal") or {}), "idRegistrasi": new["registrationId"], "namaSubPenyalur": new["name"], "alamatSubPenyalur": new["address"]}
    return out


def new_record(new: dict, new_id: str) -> dict:
    base = {
        "id": new_id, "agentId": new["agentId"], "agentName": new["agentName"],
        "status": "active", "isDeleted": False, "sourceType": "AGENT_DOCUMENT",
        "sourceDate": "2026-09-30", "sourceUrl": f"Dokumen Pangkalan/{new['documentSource']}",
        "sourceLine": new["documentRow"], "administrativeCoordinateSource": None,
        "sourceOriginal": {"idRegistrasi": new["registrationId"], "namaSubPenyalur": new["name"], "alamatSubPenyalur": new["address"], "namaPenyalur": new["agentName"]},
    }
    return update_record(base, new)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true", help="Tulis pembaruan ke master JSON")
    args = parser.parse_args()
    master = json.loads(MASTER.read_text(encoding="utf-8"))
    incoming = []
    incoming += parse_nur(DOCS / "DAFTAR PANGKALAN 2026 PT NURCAHAYA ENERGI ABADI.pdf")
    incoming += parse_wahyu(DOCS / "Daftar Pangkalan LPG 3 Kg PT. Wahyu Dwi Kencana Mandiri.pdf")
    incoming += parse_kmu(DOCS / "Daftar Pangkalan PT. KMU.xlsx")
    matches, additions = match_rows(master, incoming)
    result = deepcopy(master)
    audit = []
    for i, old, new, score in matches:
        result[i] = update_record(old, new)
        audit.append({"action": "UPDATE", "id": old["id"], "score": round(score, 4), "oldName": old["name"], **new})
    next_id = max(int(x["id"].split("-")[1]) for x in master) + 1
    for new in additions:
        new_id = f"PG-{next_id:06d}"; next_id += 1
        result.append(new_record(new, new_id))
        audit.append({"action": "ADD", "id": new_id, "score": None, "oldName": None, **new})
    matched_old_ids = {old["id"] for _, old, _, _ in matches}
    old_not_in_documents = [x["id"] for x in master if x["agentName"] in {a[1] for a in AGENTS.values()} and x["id"] not in matched_old_ids]
    summary = {
        "sourceRows": len(incoming), "updates": len(matches), "additions": len(additions),
        "oldRowsNotRepresentedInDocuments": len(old_not_in_documents),
        "oldIdsNotRepresentedInDocuments": old_not_in_documents,
        "byAgent": {agent: {"document": sum(x["agentName"] == agent for x in incoming), "updated": sum(x[2]["agentName"] == agent for x in matches), "added": sum(x["agentName"] == agent for x in additions)} for _, agent in AGENTS.values()},
    }
    REPORT_JSON.write_text(json.dumps({"summary": summary, "audit": audit}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    fields = ["action","id","score","agentName","registrationId","oldName","name","kecamatanSource","desaKelurahanSource","address","latitude","longitude","documentSource","documentRow"]
    with REPORT_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore"); writer.writeheader(); writer.writerows(audit)
    if args.apply:
        MASTER.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
