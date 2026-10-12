"""Import laporan parsial Agen 3444 untuk 1–8 Oktober 2026."""
from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"D:\# DOWNLOAD\Penyaluran Agen 3444 01-10-08-10-2026.xlsx")
MASTER_PATH = ROOT / "assets/data/lpg-pangkalan.json"
AGENTS_PATH = ROOT / "assets/data/lpg-agents.json"
OUTPUT = ROOT / "assets/data/lpg-distribution-2026-10.json"
AGENT_ID = "AG-007"
AGENT_NAME = "PT. H. AMIRUDDIN RAHMAN"
SOLD_TO = "818238"
AGENT_REPORT_CODE = "3444"


def norm(value):
    text = unicodedata.normalize("NFKD", str(value or "").upper()).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]", "", text)


def parse():
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)
    sheet = workbook["Penyaluran Agen"]
    # File ekspor menyatakan dimension=A1 meskipun memiliki 98 baris.
    sheet.reset_dimensions()
    values = list(sheet.iter_rows(values_only=True))
    headers = [str(value).strip() for value in values[0]]
    expected_dates = [f"{day:02d}-10-2026" for day in range(1, 9)]
    if headers[:3] != ["Id Pangkalan", "Id Registrasi", "Nama Pangkalan"] or headers[3:] != expected_dates:
        raise ValueError(f"Header tidak sesuai: {headers}")
    records = []
    for row_number, row in enumerate(values[1:], 2):
        daily = [int(value or 0) for value in row[3:11]]
        if any(value < 0 for value in daily):
            raise ValueError(f"Nilai negatif pada baris {row_number}")
        records.append({
            "rowNumber": row_number,
            "stationId": str(row[0]).strip(),
            "registrationId": str(row[1]).strip(),
            "reportPangkalanName": str(row[2]).strip(),
            "daily": daily,
        })
    if len(records) != 97 or len({x["stationId"] for x in records}) != 97 or len({x["registrationId"] for x in records}) != 97:
        raise ValueError("Jumlah atau keunikan laporan tidak sesuai")
    return records


def main():
    source_rows = parse()
    master = json.loads(MASTER_PATH.read_text(encoding="utf-8"))
    agents = json.loads(AGENTS_PATH.read_text(encoding="utf-8"))
    agent_rows = [row for row in master if row.get("agentId") == AGENT_ID and not row.get("isDeleted")]
    by_name = defaultdict(list)
    for row in agent_rows:
        by_name[norm(row.get("name"))].append(row)

    output_records = []
    matched_ids = set()
    for source in source_rows:
        candidates = by_name.get(norm(source["reportPangkalanName"]), [])
        if source["reportPangkalanName"] == "SUDIRMAN T":
            candidates = [row for row in agent_rows if row["id"] == "PG-000039"]
        target = candidates[0] if len(candidates) == 1 else None
        match_status = "MATCHED_AGENT_NAME"
        if source["reportPangkalanName"] == "SUDIRMAN T":
            match_status = "MATCHED_APPROVED_NAME_VARIANT"
        if source["reportPangkalanName"] == "ISMAIL":
            target = None
            match_status = "UNMATCHED_REQUIRES_ADDRESS"
        total = sum(source["daily"])
        if target:
            if target["id"] in matched_ids:
                raise ValueError(f"Target master dipakai dua kali: {target['id']}")
            matched_ids.add(target["id"])
            if source["reportPangkalanName"] == "SUDIRMAN T":
                target["previousName"] = target["name"]
                target["name"] = "SUDIRMAN T"
                target["normalizedName"] = "SUDIRMANT"
            target.update({
                "pertaminaStationId": source["stationId"],
                "registrationId": source["registrationId"],
                "registrationVerificationStatus": "VERIFIED_IN_PERTAMINA_DISTRIBUTION_REPORT",
                "registrationVerificationSource": SOURCE.name,
                "latestDistributionPeriod": "2026-10-01/2026-10-08",
                "latestDistributionNormal": total,
                "latestDistributionFacultative": 0,
                "latestDistributionTotal": total,
                "latestDistributionUnit": "tabung",
                "latestDistributionSource": SOURCE.name,
                "dataVersion": "2026-10-09-agent-3444-partial-import-v1",
            })
            if target["id"] == "PG-000482":
                target.update({
                    "latestIdentityEvidence": "PERTAMINA_DISTRIBUTION_REPORT_2026-10-01_08",
                    "operationalDataStatus": "CONFIRMED_IN_LATEST_AVAILABLE_AGENT_REPORT",
                    "identityReviewNote": "Bukan duplikat PG-000478. Laporan Agen 3444 tanggal 1–8 Oktober 2026 mengonfirmasi record AG-007 ini secara terpisah dengan ID pangkalan Pertamina 400581 dan registrasi 791272818238104.",
                    "identityReviewDate": "2026-10-09",
                })
        output_records.append({
            "period": "2026-10", "periodStart": "2026-10-01", "periodEnd": "2026-10-08", "unit": "tabung",
            "agentId": AGENT_ID, "agentName": AGENT_NAME, "agentReportCode": AGENT_REPORT_CODE, "soldTo": SOLD_TO,
            "stationId": source["stationId"], "registrationId": source["registrationId"],
            "pangkalanId": target["id"] if target else None,
            "pangkalanName": target["name"] if target else source["reportPangkalanName"],
            "reportPangkalanName": source["reportPangkalanName"],
            "dailyDistribution": {f"2026-10-{day:02d}": value for day, value in enumerate(source["daily"], 1)},
            "totalNormal": total, "totalFacultative": 0, "grandTotal": total,
            "sourceDocument": SOURCE.name, "sourceRow": source["rowNumber"],
            "sourceStatus": "AGENT_REPORTED_PERTAMINA_SYSTEM", "attestation": "SYSTEM_EXPORT_UNSIGNED",
            "masterMatchStatus": match_status,
        })

    totals = {f"2026-10-{day:02d}": sum(row["dailyDistribution"][f"2026-10-{day:02d}"] for row in output_records) for day in range(1, 9)}
    report = {
        "schemaVersion": 1, "period": "2026-10", "periodStart": "2026-10-01", "periodEnd": "2026-10-08",
        "unit": "tabung", "status": "agent-reported-partial-period", "generatedAt": "2026-10-09T17:47:04+08:00",
        "coverage": {"reportedAgentCount": 1, "officialAgentCount": 9, "isCompleteRegencyCoverage": False,
                     "isCompleteMonthlyPeriod": False, "warning": "Hanya Agen 3444 untuk 1–8 Oktober 2026; bukan total bulanan atau total Kabupaten Pinrang."},
        "recordCount": len(output_records),
        "totals": {"totalNormal": sum(totals.values()), "totalFacultative": 0, "grandTotal": sum(totals.values()),
                   "kilograms": sum(totals.values()) * 3, "metricTons": round(sum(totals.values()) * 3 / 1000, 3)},
        "dataQuality": {"arithmeticValidatedRecordCount": 97, "uniqueStationIdCount": 97, "uniqueRegistrationCount": 97,
                        "masterMatchedRecordCount": 96, "unmatchedIdentityCount": 1, "duplicateRegistrationCount": 0,
                        "gpsFieldVerifiedCount": 0},
        "dailyTotals": totals,
        "summaries": [{"agentId": AGENT_ID, "agentName": AGENT_NAME, "agentReportCode": AGENT_REPORT_CODE, "soldTo": SOLD_TO,
                       "pangkalanCount": 97, "masterMatchedCount": 96, "zeroDistributionCount": sum(x["grandTotal"] == 0 for x in output_records),
                       "grandTotal": sum(totals.values()), "unit": "tabung", "sourceDocument": SOURCE.name}],
        "unmatchedRecords": [row for row in output_records if not row["pangkalanId"]],
        "records": output_records,
    }
    if report["totals"]["grandTotal"] != 10640 or totals["2026-10-04"] != 0:
        raise ValueError("Rekap laporan tidak cocok")
    OUTPUT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    MASTER_PATH.write_text(json.dumps(master, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    agent = next(row for row in agents if row["id"] == AGENT_ID)
    agent.update({"soldTo": SOLD_TO, "agentReportCode": AGENT_REPORT_CODE, "linkedPangkalanCount": len(agent_rows),
                  "reportedPangkalanCount": 97, "matchedReportedPangkalanCount": 96,
                  "latestDistributionPeriod": "2026-10-01/2026-10-08", "latestDistributionTotal": 10640,
                  "distributionDataUnit": "tabung", "distributionDataSource": SOURCE.name,
                  "relationshipStatus": "96 dari 97 baris laporan Oktober terhubung; ISMAIL menunggu alamat untuk pencocokan."})
    AGENTS_PATH.write_text(json.dumps(agents, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"records": 97, "matched": 96, "unmatched": 1, "grandTotal": 10640,
                      "matchStatus": Counter(row["masterMatchStatus"] for row in output_records)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
